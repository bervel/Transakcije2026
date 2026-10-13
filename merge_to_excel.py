# -*- coding: utf-8 -*-
import re, os, sys, json, shutil
os.chdir(r"d:\BVDoc\Transakcije2026")
sys.stdout.reconfigure(encoding="utf-8")
from datetime import date, timedelta
import openpyxl

# backup
if not os.path.exists(r"Transakcije (5) - backup.xlsx"):
    shutil.copy(r"Transakcije (5).xlsx", r"Transakcije (5) - backup.xlsx")

items=json.load(open("toadd.json",encoding="utf-8"))
def fd(d): return f"{d.day}. {d.month}. {d.year}"

rows=[]
for it in items:
    d=date.fromisoformat(it["d"])
    v=it["v"]; cur=it["cur"]
    txt=f"NAKUP {it['desc']} {v:,.2f} {cur}".replace(","," ")
    sign="-" if not it["neg"] else "+"
    rows.append([txt,it["cat"],sign,round(v,2),cur,fd(d),"","","","","","","",fd(d+timedelta(days=2)),"",""])

wb=openpyxl.load_workbook(r"Transakcije (5).xlsx")
ws=wb["transakcije"]

def pd(v):
    p=[int(x) for x in re.findall(r"\d+",str(v))]
    return date(p[2],p[1],p[0]) if len(p)>=3 else None

# sheet is newest-first (descending). Insert oldest→newest (ascending date) order, each time
# finding the first row with date strictly < newd and inserting before it. Already-inserted
# (older) rows sit below and won't be encountered first; equal dates preserve grouping.
def find_insert(newd):
    for i in range(2, ws.max_row+1):
        d=pd(ws.cell(i,6).value)
        if d and d<newd:
            return i
    return ws.max_row+1

import unicodedata
print("rows to insert:",len(rows))
indexed=enumerate(items)
newdates=sorted(((i,date.fromisoformat(it["d"]),it) for i,it in indexed), key=lambda x: x[1])
inserted=0
for _i,nd,it in newdates:
    idx=find_insert(nd)
    v=it["v"]; cur=it["cur"]
    txt=f"NAKUP {it['desc']} {v:,.2f} {cur}".replace(","," ")
    sign="-" if not it["neg"] else "+"
    row=[txt,it["cat"],sign,round(v,2),cur,fd(nd),"","","","","","","",fd(nd+timedelta(days=2)),"",""]
    ws.insert_rows(idx,1)
    for c,val in enumerate(row,start=1):
        ws.cell(idx,c).value=val
    inserted+=1
    if inserted%50==0: print(inserted,"done")
wb.save(r"Transakcije (5).xlsx")
print("inserted:",inserted)
wb=openpyxl.load_workbook(r"Transakcije (5).xlsx")
ws=wb["transakcije"]
print("new max_row:",ws.max_row)
# verify sort
bad=0; prev=None
for i in range(2,ws.max_row+1):
    d=pd(ws.cell(i,6).value)
    if d and prev and d>prev: bad+=1; print("sort break at row",i,prev,d)
    if d: prev=d
print("sort violations:",bad)
